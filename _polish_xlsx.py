# -*- coding: utf-8 -*-
"""核查表版式收尾：11 列对象 + 第 14 列标注后，补齐列宽与跨列合并。

只在结构（列数/行数）已确定后运行：
  - 对象列 C~M 统一宽度（表头「华为交换机」等 5 字需 ≥11）
  - 标题行 A1:N1、备注行 B140:N140 与表宽对齐
"""
import openpyxl

XLSX = '配置核查表_v2.0.0_标注自动验证.xlsx'
wb = openpyxl.load_workbook(XLSX)
ws = wb['Sheet1']
assert ws.max_column == 14, ws.max_column

# 1) 对象列宽统一（C~M）
for col in 'CDEFGHIJKLM':
    ws.column_dimensions[col].width = 11.0

# 2) 跨列合并对齐表宽（先解除再重建；被合并区域内的单元格需为空）
for rng, target in (('A1:L1', 'A1:N1'), ('B140:L140', 'B140:N140')):
    if rng in [str(m) for m in ws.merged_cells.ranges]:
        ws.unmerge_cells(rng)
    tail = target.split(':')[1]
    tcol = openpyxl.utils.column_index_from_string(''.join(c for c in tail if c.isalpha()))
    for c in range(2, tcol + 1):
        assert ws.cell(int(target[1:]) if target[1:].isdigit() else 1, c).value in (None, '') or c <= 12, \
            f'{target} 范围内有值，需先确认'
    ws.merge_cells(target)

wb.save(XLSX)
print('版式收尾完成')
