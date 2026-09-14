# -*- coding: utf-8 -*-
r"""第 3、4 章无图小节配图生成并插入（58 节，每节 1 图）。

生成 _tmp_figs34/g*.png（与既有指导书配图同风格：白底黑字等宽终端渲染），
再按 dump 段落索引锚定各节末段之后插入（居中、宽度同既有图规则）。

用法：
  python _gen_ch34_figs.py            # 只生成图片
  python _gen_ch34_figs.py --apply    # 插入 docx（需先备份）
"""
import os
import subprocess
import sys

from PIL import Image, ImageChops
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm

HERE = os.path.dirname(os.path.abspath(__file__))
FIGS = os.path.join(HERE, "_tmp_figs34")
PATH = os.path.join(HERE, "配置核查作业指导书_v2.2.docx")
CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
APPLY = "--apply" in sys.argv
DPI = 110.0


def shoot(html_path, png_path, w, h):
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-sandbox",
                    "--hide-scrollbars", "--force-device-scale-factor=1",
                    "--virtual-time-budget=3000", "--window-size=%d,%d" % (w, h),
                    "--screenshot=" + png_path, "file:///" + html_path.replace("\\", "/")],
                   capture_output=True, text=True)


def save_trim(png_path, pad=8):
    im = Image.open(png_path).convert("RGB")
    bg = Image.new("RGB", im.size, (255, 255, 255))
    bbox = ImageChops.difference(im, bg).getbbox()
    if bbox:
        im.crop((max(bbox[0] - pad, 0), max(bbox[1] - pad, 0),
                 min(bbox[2] + pad, im.width), min(bbox[3] + pad, im.height))).save(png_path)
    return Image.open(png_path).size


def term(name, lines):
    body = "\n".join(lines).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    html_path = os.path.join(FIGS, "_t.html")
    png_path = os.path.join(FIGS, name + ".png")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write('<!DOCTYPE html><html><head><meta charset="UTF-8"><style>'
                'body{margin:0;background:#fff;}'
                'pre{font:15px/1.9 "Courier New","Consolas",monospace;color:#000;'
                'margin:0;padding:6px 10px;white-space:pre;}</style></head>'
                '<body><pre>%s</pre></body></html>' % body)
    shoot(html_path, png_path, 1150, 900)
    if os.path.exists(html_path):
        os.remove(html_path)
    print(name, save_trim(png_path))


# ============ 58 张图定义：(png 名, 段落索引, 锚文本前缀校验, 终端行) ============
FIGSPEC = [
    # ---------- 第 3 章 ----------
    ("g313", 4188, "【判定】数据库层（TDE 或列级加密）", [
        "[root@kylin-server ~]# mysql -uroot -p -e \"SHOW VARIABLES LIKE 'default_table_encryption'\"",
        "+--------------------------+-------+",
        "| Variable_name            | Value |",
        "+--------------------------+-------+",
        "| default_table_encryption | ON    |",
        "+--------------------------+-------+",
        "[root@kylin-server ~]# mysql -uroot -p -e \"SHOW PLUGINS\" | grep -i keyring",
        "keyring_file         ACTIVE   FILE       keyring_file.so    GPL",
        "[root@kylin-server ~]# grep -r \"encrypt\\|aes\\|sm4\" /opt/app/conf/*.conf",
        "app.conf: data.field.cipher = SM4",
        "app.conf: data.key.store = /etc/app/keys/master.ks",
    ]),
    ("g323", 4216, "【判定】文件传输/交互经统一平台", [
        "[root@kylin-server ~]# grep -r \"transfer\\|sftp\\|ftp\" /opt/*/conf/*.conf | head -2",
        "/opt/ferry/conf/app.conf: transfer.channel = unify-fs",
        "/opt/ferry/conf/app.conf: transfer.whitelist = 10.30.8.0/24",
        "[root@kylin-server ~]# ss -tlnp | grep -E ':21|:22|:445'",
        "LISTEN 0 128 10.30.8.11:445 users:((\"smbd\",pid=2213,fd=35))",
        "LISTEN 0 128 0.0.0.0:22     users:((\"sshd\",pid=1180,fd=3))",
        "[root@kylin-server ~]# tail -2 /var/log/samba/log.smbd",
        "[2026/09/02 14:21:07] zhang_san 10.30.8.15 put design_v2.docx (2.3MB) audit_id=8841",
        "[2026/09/02 15:02:44] li_si    10.30.8.22 get design_v2.docx (2.3MB) audit_id=8856",
    ]),
    ("g333", 4237, "【判定】使用专业工具按标准算法", [
        "[root@kylin-server ~]# which dban nwipe wipe",
        "/usr/bin/nwipe",
        "/usr/bin/wipe",
        "[root@kylin-server ~]# hdparm --user-master u --security-set-pass NULL /dev/sdb",
        "security_set_password: \"NULL\" issued",
        "[root@kylin-server ~]# hdparm -I /dev/sdb | grep -i erase",
        "\tsupported: enhanced erase",
        "\t32min for SECURITY ERASE UNIT.",
    ]),
    ("g341", 4246, "【判定】按载体类型采用对应物理销毁", [
        r'C:\Windows\system32> dir /s /b "D:\保密管理\载体销毁记录"',
        r'D:\保密管理\载体销毁记录\2026-03 消磁记录-移动硬盘SN8K21.xlsx',
        r'D:\保密管理\载体销毁记录\2026-03 销毁监督签认表-0312.pdf',
        r'D:\保密管理\载体销毁记录\2026-05 粉碎销毁凭证-U盘SN3377.pdf',
        r'D:\保密管理\载体销毁记录\2026-08 光盘粉碎记录-0821.xlsx',
        r'D:\保密管理\载体销毁记录\销毁影像资料\20260821_1015.jpg',
        "",
        r'C:\Windows\system32>',
    ]),
    ("g353", 4289, "【判定】系统日志、应用日志、数据库审计", [
        "[root@kylin-server ~]# ls -la /var/log/tomcat/ | head -3",
        "drwxrwx---  2 tomcat tomcat  4096 9月 12 08:30 .",
        "-rw-r-----  1 tomcat tomcat   31M 9月 12 08:30 catalina.out",
        "[root@kylin-server ~]# mysql -uroot -p -e \"SHOW VARIABLES LIKE 'audit%'\"",
        "Empty set (0.00 sec)",
        "# 社区版无audit插件：需结合 general_log/binlog 及企业审计平台确认",
        "[root@kylin-server ~]# grep -r \"tls\\|encrypt\" /etc/rsyslog.conf",
        "*.* @@(o)10.30.1.55:6514;RSYSLOG_Syslog_PROTOCOL23Format  # TLS加密转发",
    ]),
    ("g361", 4296, "如现场部署了 IPS、WAF", [
        "<FW-OUT-01> display security-policy rule name DATA-OUT",
        " rule name DATA-OUT",
        "  source-zone trust",
        "  destination-zone untrust",
        "  source-address 10.30.8.0 mask 255.255.255.0",
        "  service https ftp",
        "  action permit",
        "  content-filter enable",
        "<FW-OUT-01> display security-policy rule name BLOCK-UPLOAD",
        " rule name BLOCK-UPLOAD",
        "  source-zone trust",
        "  destination-zone untrust",
        "  service http ftp",
        "  action deny",
    ]),
    ("g362", 4300, "若抽查发现边界仅放行必要业务", [
        "<FW-OUT-01> display content-filter log brief | include deny",
        "2026-09-11 10:22:31  10.30.8.15 -> 203.0.113.9  HTTP-POST  关键字命中:机密   deny",
        "2026-09-11 11:05:12  10.30.8.22 -> 198.51.100.7 FTP-STOR  文件类型:*.dwg    deny",
        "2026-09-11 14:48:03  10.30.8.31 -> 203.0.113.9  SMTP-ATT   正则命中:身份证号 alert",
        "<FW-OUT-01>",
    ]),
    ("g363", 4305, "【判定】已部署 DLP", [
        "[root@kylin-server ~]# ps aux | grep -i dlp | grep -v grep",
        "root   2811  0.8  1.2 412388 24560 ?  Ssl  08:30  1:12 /opt/dlpc/bin/dlpagent",
        "root   2813  0.2  0.8 305212 18304 ?  Ssl  08:30  0:31 /opt/dlpc/bin/dlpwatch",
    ]),
    ("g373", 4331, "【判定】数据库管理登录具备增强认证", [
        "mysql> SELECT user,host,plugin FROM mysql.user WHERE user IN ('root','admin','test');",
        "+---------+-----------+-----------------------+",
        "| user    | host      | plugin                |",
        "+---------+-----------+-----------------------+",
        "| root    | localhost | caching_sha2_password |",
        "+---------+-----------+-----------------------+",
        "mysql> SELECT user,host FROM mysql.user WHERE authentication_string='';",
        "Empty set (0.00 sec)",
        "mysql> SHOW VARIABLES LIKE 'general_log';",
        "+---------------+-------+",
        "| Variable_name | Value |",
        "+---------------+-------+",
        "| general_log   | ON    |",
        "+---------------+-------+",
    ]),
    ("g383", 4358, "【判定】数据库按重要程度划分", [
        "mysql> SHOW DATABASES;",
        "+--------------------+",
        "| Database           |",
        "+--------------------+",
        "| information_schema |",
        "| mysql              |",
        "| pubdb              |",
        "| secdb              |",
        "| sys                |",
        "+--------------------+",
        "mysql> SHOW GRANTS FOR 'sec_user'@'%';",
        "+----------------------------------------------------+",
        "| GRANT SELECT,INSERT ON `secdb`.* TO `sec_user`@`%` |",
        "+----------------------------------------------------+",
    ]),
    ("g393", 4394, "【判定】应用层 HTTPS/TLS", [
        "[root@kylin-server ~]# openssl s_client -connect localhost:443 2>/dev/null | grep -i \"protocol\\|cipher\"",
        "New, TLSv1.3, Cipher is TLS_AES_256_GCM_SHA384",
        "[root@kylin-server ~]# mysql -uroot -p -e \"SHOW STATUS LIKE 'Ssl_cipher'\"",
        "+---------------+-----------------------------+",
        "| Variable_name | Value                       |",
        "+---------------+-----------------------------+",
        "| Ssl_cipher    | ECDHE-RSA-AES128-GCM-SHA256 |",
        "+---------------+-----------------------------+",
        "[root@kylin-server ~]# mysql -uroot -p -e \"SHOW VARIABLES LIKE 'require_secure_transport'\"",
        "+--------------------------+-------+",
        "| require_secure_transport | ON    |",
        "+--------------------------+-------+",
    ]),
    ("g3103", 4433, "【判定】数据共享经统一平台", [
        "[root@kylin-server ~]# ss -tlnp | grep -E ':139|:445|:2049|:21|:22'",
        "LISTEN 0 50 0.0.0.0:445 users:((\"smbd\",pid=2213,fd=34))",
        "LISTEN 0 50 0.0.0.0:139 users:((\"smbd\",pid=2213,fd=33))",
        "LISTEN 0 128 0.0.0.0:22  users:((\"sshd\",pid=1180,fd=3))",
        "[root@kylin-server ~]# smbstatus -S",
        "Service   machine     Connect at",
        "share01   10.30.8.15  九 09 10:15:22",
        "[root@kylin-server ~]# tail -1 /var/log/samba/log.smbd",
        "[2026/09/09 10:16:40] zhang_san 10.30.8.15 create audit_2026.docx (1.1MB)",
    ]),
    ("g3113", 4457, "【判定】数据访问权限分级", [
        "mysql> SHOW GRANTS FOR 'sec_user'@'%';",
        "+---------------------------------------------------------------+",
        "| GRANT SELECT,INSERT ON `secdb`.`secret_tbl` TO `sec_user`@`%` |",
        "+---------------------------------------------------------------+",
        "[root@kylin-server ~]# mysql -uroot -p -e \"SELECT GRANTEE,TABLE_NAME,COLUMN_NAME,PRIVILEGE_TYPE FROM information_schema.COLUMN_PRIVILEGES LIMIT 3\"",
        "'sec_user'@'%'  secret_tbl  content  SELECT",
        "'sec_user'@'%'  secret_tbl  level    SELECT",
        "[root@kylin-server ~]# grep -r \"role\\|权限\" /opt/app/conf/*.conf",
        "app.conf: app.role.mappings = /etc/app/role-map.xml",
    ]),
    ("g3123", 4474, "【判定】采集范围与业务需求", [
        "[root@kylin-server ~]# grep -r \"collect\\|字段\" /opt/app/conf/*.conf",
        "app.conf: collect.fields = uid,op_type,op_time",
        "app.conf: collect.excludes = gps,contacts,sensor",
        "[root@kylin-server ~]# tail -2 /var/log/app/collect.log",
        "2026-09-12 09:14:01 src=app-ui uid=2041 op=search freq=3/min",
        "2026-09-12 09:14:22 src=app-ui uid=2041 op=query  freq=5/min",
    ]),
    ("g3133", 4504, "【判定】数据在存储、处理、传输", [
        "[root@kylin-server ~]# grep -r \"密级\\|classified\" /opt/app/conf/*.conf",
        "app.conf: app.classified.label = on",
        "app.conf: app.classified.levels = 公开,内部,秘密,机密",
        "[root@kylin-server ~]# mysql -uroot -p -e \"SELECT column_name FROM information_schema.columns WHERE table_schema='secdb' AND table_name='doc_info'\"",
        "+-------------+",
        "| column_name |",
        "+-------------+",
        "| doc_id      |",
        "| title       |",
        "| level_code  |",
        "| content     |",
        "+-------------+",
    ]),
    ("g3143", 4537, "【判定】大数据访问需通过", [
        "[hdfs@bigdata-01 ~]$ hdfs dfs -getfacl /data/classified",
        "# file: /data/classified",
        "# owner: hive   # group: hive",
        "user::rwx",
        "user:sec_auditor:r-x",
        "group::r-x",
        "other::---",
        "[hdfs@bigdata-01 ~]$ klist",
        "Default principal: hive@BIGDATA.LOCAL",
        "09/12/2026 08:31:02  09/13/2026 08:31:02  krbtgt/BIGDATA.LOCAL@BIGDATA.LOCAL",
        "[hdfs@bigdata-01 ~]$ ss -tlnp | grep -E ':50070|:10000|:8020'",
        "LISTEN 0 50 0.0.0.0:10000 users:((\"java\",pid=3311,fd=180))",
        "LISTEN 0 128 0.0.0.0:8020 users:((\"java\",pid=2877,fd=221))",
    ]),
    # ---------- 第 4 章 ----------
    ("g413", 4565, "【判定】应用层、数据库层、操作系统层", [
        "[root@kylin-server ~]# crontab -l | grep mysqldump",
        "0 2 * * * /opt/scripts/backup_mysql.sh",
        "[root@kylin-server ~]# mysql -uroot -p -e \"SHOW VARIABLES LIKE 'log_bin'\"",
        "+---------------+-------+",
        "| log_bin       | ON    |",
        "+---------------+-------+",
        "[root@kylin-server ~]# find /opt -name \"*restore*\" 2>/dev/null",
        "/opt/scripts/restore_mysql.sh",
    ]),
    ("g423", 4581, "【判定】操作系统、应用、数据库任一层", [
        "mysql> SELECT VERSION();",
        "+-----------+",
        "| VERSION() |",
        "+-----------+",
        "| 8.0.36    |",
        "+-----------+",
        "[root@kylin-server ~]# rpm -qa | grep -i app-server",
        "app-server-3.2.7-1.ky10.x86_64",
    ]),
    ("g433", 4604, "【判定】应用基于可信根签名", [
        "[root@kylin-server ~]# rpm -K /opt/pkg/app-server-3.2.7-1.ky10.x86_64.rpm",
        "/opt/pkg/app-server-3.2.7-1.ky10.x86_64.rpm: rsa sha1 (md5) PGP md5 确定",
        "[root@kylin-server ~]# systemctl list-units | grep -iE \"aide|tripwire\"",
        "aide-check.timer   loaded active waiting   AIDE 定期完整性校验",
        "[root@kylin-server ~]# systemctl status aide-check.timer | head -2",
        "Active: active (waiting) since 三 2026-09-03 08:00:11 CST",
    ]),
    ("g443", 4622, "【判定】公共服务与涉密服务", [
        "[root@kylin-server ~]# sudo ip route show",
        "default via 10.30.8.254 dev eth0",
        "10.30.8.0/24 dev eth0 proto kernel scope link src 10.30.8.11",
        "192.168.50.0/24 via 10.30.8.253 dev eth1",
        "[root@kylin-server ~]# sudo iptables -L INPUT -n | head -4",
        "Chain INPUT (policy DROP)",
        "ACCEPT  tcp  --  10.30.8.0/24  0.0.0.0/0  tcp dpt:443",
        "DROP    all  --  0.0.0.0/0      0.0.0.0/0",
    ]),
    ("g453", 4642, "【判定】防DDoS能力需主机", [
        "[root@kylin-server ~]# grep -rE \"limit_req|limit_conn\" /etc/nginx/ | grep -v \"#\"",
        "nginx.conf:      limit_req_zone $binary_remote_addr zone=one:10m rate=10r/s;",
        "conf.d/app.conf: limit_req zone=one burst=20 nodelay;",
        "[root@kylin-server ~]# grep -E \"maxThreads|acceptCount\" /opt/tomcat/conf/server.xml",
        "        maxThreads=\"200\" acceptCount=\"50\"",
    ]),
    ("g461", 4652, "附加防护措施核查。", [
        "[root@kylin-server ~]# curl -s \"http://localhost/search?kw=<script>alert(1)</script>\" | grep -o \"&lt;script&gt;[^<]*\"",
        "&lt;script&gt;alert(1)&lt;/script&gt;",
        "# 输出已被 HTML 编码转义，脚本未被执行",
        "[root@kylin-server ~]# curl -s \"http://localhost/login?u=admin'--\" | tail -1",
        "{\"code\":400,\"msg\":\"请求参数不合法，请联系管理员\"}",
        "# 未返回数据库报错、SQL 语句等敏感信息",
    ]),
    ("g462", 4659, "若结合实际测试发现典型 SQL 注入", [
        "[root@kylin-server ~]# grep -c \"deny\\|拦截\" /var/log/app/waf_audit.log",
        "1284",
        "[root@kylin-server ~]# tail -3 /var/log/app/waf_audit.log",
        "2026-09-12 10:02:11  rule=SQLi-001  src=10.30.9.7   uri=/login  action=deny",
        "2026-09-12 10:14:38  rule=XSS-014   src=10.30.9.19  uri=/search action=encode",
        "2026-09-12 11:01:52  rule=SQLi-003  src=10.30.9.7   uri=/query  action=deny",
    ]),
    ("g463", 4662, "【判定】应用数据库连接账号必须", [
        "mysql> SELECT user,host FROM mysql.user;",
        "+------------------+-----------+",
        "| user             | host      |",
        "+------------------+-----------+",
        "| appuser          | %         |",
        "| mysql.infoschema | localhost |",
        "| mysql.session    | localhost |",
        "| root             | localhost |",
        "+------------------+-----------+",
        "mysql> SHOW GRANTS FOR 'appuser'@'%';",
        "+----------------------------------------------------------+",
        "| GRANT SELECT,INSERT,UPDATE ON `appdb`.* TO `appuser`@`%` |",
        "+----------------------------------------------------------+",
    ]),
    ("g473", 4680, "【判定】网站服务应仅发布静态页面", [
        "[root@kylin-server ~]# grep -rE \"fastcgi_pass|\\.php|\\.jsp\" /etc/nginx/nginx.conf /etc/nginx/conf.d/*.conf 2>/dev/null",
        "# grep 无输出：未配置 fastcgi_pass/动态脚本解析",
        "[root@kylin-server ~]# ls -la /opt/tomcat/webapps/",
        "drwxr-x---  3 tomcat tomcat 4096 6月 30 10:12 ROOT",
        "-rw-r-----  1 tomcat tomcat 5213 6月 30 10:12 index.html",
    ]),
    ("g481", 4687, "结合网站页面展示形式", [
        r"C:\Windows\system32> dir D:\wwwroot",
        r" D:\wwwroot 的目录",
        "",
        "2026/06/30  10:12    <DIR>          css",
        "2026/06/30  10:12    <DIR>          images",
        "2026/06/30  10:12    <DIR>          js",
        "2026/06/30  10:12             5,213 index.html",
        "2026/06/30  10:12             8,942 news.html",
        "               2 个文件         14,155 字节",
        r"C:\Windows\system32> netstat -ano | findstr LISTENING | findstr \":80 :443\"",
        "  TCP    0.0.0.0:80     0.0.0.0:0      LISTENING    1980",
        "  TCP    0.0.0.0:443    0.0.0.0:0      LISTENING    1980",
    ]),
    ("g483", 4699, "【判定】部署专业防篡改系统", [
        "[root@kylin-server ~]# systemctl list-units | grep -iE \"guard|tamper|protect\"",
        "tamper-guard.service   loaded active running   网页防篡改守护进程",
        "[root@kylin-server ~]# grep -E \"/var/www|html\" /etc/aide/aide.conf",
        "/var/www/html CONTENT_EX",
        "[root@kylin-server ~]# mount | grep \"/var/www\"",
        "/dev/mapper/vg-www on /var/www type ext4 (ro,bind)",
    ]),
    ("g491", 4706, "综合判定。上述各项核查结果均符合", [
        "[root@test-client ~]# curl -s \"http://10.30.8.21/login\" -d \"user=admin&pwd=' or '1'='1\"",
        "{\"code\":401,\"msg\":\"用户名或口令错误\"}",
        "# 未返回数据库错误、未绕过认证",
        "[root@test-client ~]# curl -s \"http://10.30.8.21/search?kw=%3Cscript%3Ealert(1)%3C/script%3E\" | grep -o \"kw=[^<]*\"",
        "kw=&lt;script&gt;alert(1)&lt;/script&gt;",
        "# 注入脚本被转义，未执行",
    ]),
    ("g493", 4714, "【判定】应用需在代码层使用参数化", [
        "[root@kylin-server ~]# grep -rn \"Statement\\|concat\\|拼接\" /opt/app/src/ 2>/dev/null | head -2",
        "/opt/app/src/dao/UserDao.java:  PreparedStatement ps = conn.prepareStatement(SQL);",
        "/opt/app/src/dao/DocDao.java:  PreparedStatement ps = conn.prepareStatement(SQL);",
        "[root@kylin-server ~]# grep -r \"rule id\" /opt/app/conf/waf.xml",
        "<waf:rule id=\"SQLi-001\" action=\"deny\"   enable=\"true\"/>",
        "<waf:rule id=\"XSS-014\"  action=\"encode\" enable=\"true\"/>",
    ]),
    ("g4101", 4721, "结合系统配置、功能测试结果", [
        "[root@kylin-server ~]# curl -s -o /dev/null -w \"%{http_code}\" -F \"file=@shell.jsp\" http://localhost/upload",
        "403",
        "# 可执行文件上传被拒绝",
        "[root@kylin-server ~]# grep -r \"upload\" /opt/app/conf/app.conf",
        "app.conf: app.upload.whitelist = jpg,jpeg,png,pdf,docx",
        "app.conf: app.upload.exec-deny = true",
    ]),
    ("g4103", 4731, "【判定】应用需对可执行代码", [
        "[root@kylin-server ~]# grep -rnE \"eval|exec|system|shell_exec\" /opt/app/ 2>/dev/null | head -2",
        "/opt/app/src/util/ExprGuard.java:  // 白名单表达式引擎，eval 已禁用",
        "/opt/app/src/main/resources/blacklist.yaml:  denied: [eval, exec, system, shell_exec]",
    ]),
    ("g4111", 4737, "判定标准：系统中的用户账户", [
        "PS C:\\Windows\\system32> Get-LocalUser | Select-Object Name,Enabled",
        "",
        "Name          Enabled",
        "----          -------",
        "Administrator  True",
        "Guest          False",
        "zhang_wei      True",
        "li_qiang       True",
        "",
        "PS C:\\Windows\\system32> Get-LocalGroupMember -Group \"Administrators\"",
        "",
        "Name           PrincipalSource",
        "----           ---------------",
        "Administrator  Local",
        "zhang_wei      Local",
    ]),
    ("g4112", 4743, "判定标准：系统用户身份", [
        "[root@kylin-server ~]# id secops",
        "uid=1001(secops) gid=1001(secops) 组=1001(secops),10(wheel)",
        "[root@kylin-server ~]# groups secops",
        "secops : secops wheel",
        "[root@kylin-server ~]# sudo -l -U secops | head -3",
        "匹配此主机上 %10.30.8 的默认条目：",
        "    (root) NOPASSWD: /usr/bin/systemctl status",
    ]),
    ("g4113", 4748, "【判定】应用自身实现RBAC", [
        "[root@kylin-server ~]# grep -r \"role\\|permission\" /opt/app/conf/*.conf | head -2",
        "/opt/app/conf/security.xml: <role name=\"secadmin\" perms=\"user:manage,audit:read\"/>",
        "/opt/app/conf/security.xml: <role name=\"audit\"    perms=\"audit:read,audit:export\"/>",
        "mysql> SELECT user,host FROM mysql.user;",
        "+----------+-----------+",
        "| user     | host      |",
        "+----------+-----------+",
        "| appuser  | %         |",
        "| root     | localhost |",
        "+----------+-----------+",
    ]),
    ("g4123", 4769, "【判定】应用会话与Web服务器", [
        "[root@kylin-server ~]# grep -r \"maxSession\\|maxConnection\" /opt/*/conf/*.conf 2>/dev/null",
        "/opt/app/conf/server.yml: max-sessions: 500",
        "/opt/app/conf/server.yml: per-user-sessions: 5",
        "[root@kylin-server ~]# grep -r \"limit_conn\\|limit_req\" /etc/nginx/ | grep -v \"#\"",
        "nginx.conf:      limit_conn_zone $binary_remote_addr zone=addr:10m;",
        "conf.d/app.conf: limit_conn addr 100;",
    ]),
    ("g4133", 4778, "【判定】业务管理终端需用途专用", [
        "[root@biz-mgmt-01 ~]# rpm -qa | grep -iE \"office|game|video\"",
        "# 无输出：未安装办公/娱乐等无关软件",
        "[root@biz-mgmt-01 ~]# grep -E \"AllowUsers\" /etc/ssh/sshd_config",
        "AllowUsers sysadmin secops",
        "[root@biz-mgmt-01 ~]# cat /etc/passwd | grep -E \"sh$\" | cut -d: -f1",
        "root",
        "sysadmin",
        "secops",
    ]),
    ("g4143", 4808, "【判定】访问控制细粒度", [
        "mysql> SHOW GRANTS FOR 'sec_writer'@'%';",
        "+--------------------------------------------------------------------+",
        "| GRANT SELECT,INSERT ON `secdb`.`secret_tbl` TO `sec_writer`@`%`    |",
        "+--------------------------------------------------------------------+",
        "[root@kylin-server ~]# grep -r \"role\\|映射\" /opt/app/conf/*.conf | head -2",
        "/opt/app/conf/perm-map.xml: <perm role=\"sec_writer\" table=\"secret_tbl\" ops=\"select,insert\"/>",
        "/opt/app/conf/perm-map.xml: <perm role=\"pub_reader\" table=\"public_tbl\" ops=\"select\"/>",
    ]),
    ("g4153", 4830, "【判定】文档专用业务系统", [
        "[root@doc-sys-01 ~]# grep -r \"signature\\|密级\\|classified\" /opt/docsapp/conf/*.conf | head -3",
        "/opt/docsapp/conf/app.conf: sign.verify = on",
        "/opt/docsapp/conf/app.conf: sign.invalid.block = true",
        "/opt/docsapp/conf/app.conf: doc.level.label = 公开,内部,秘密,机密",
        "[root@doc-sys-01 ~]# ls /data/docs/ | head -3",
        "2026-A-0175_内部_项目周报.pdf",
        "2026-M-0331_秘密_运维手册v3.docx",
        "2026-S-1102_机密_预案汇编.pdf",
    ]),
    ("g4163", 4857, "【判定】应用管理界面与数据库", [
        "[root@kylin-server ~]# openssl s_client -connect 10.30.8.21:443 2>/dev/null | grep -i \"cipher\" | head -1",
        "New, TLSv1.3, Cipher is TLS_AES_256_GCM_SHA384",
        "mysql> SHOW VARIABLES LIKE 'have_ssl';",
        "+---------------+-------+",
        "| Variable_name | Value |",
        "+---------------+-------+",
        "| have_ssl      | YES   |",
        "+---------------+-------+",
    ]),
    ("g4173", 4894, "【判定】三权分立需覆盖", [
        "mysql> SELECT user,host FROM mysql.user;",
        "+----------+-----------+",
        "| user     | host      |",
        "+----------+-----------+",
        "| appuser  | %         |",
        "| audit    | localhost |",
        "| root     | localhost |",
        "+----------+-----------+",
        "mysql> SELECT user,Super_priv FROM mysql.user WHERE user IN ('root','audit');",
        "+-------+------------+",
        "| user  | Super_priv |",
        "+-------+------------+",
        "| root  | Y          |",
        "| audit | N          |",
        "+-------+------------+",
    ]),
    ("g4183", 4917, "【判定】业务管理终端登录需在", [
        "[root@kylin-server ~]# sudo iptables -L INPUT -n --line-numbers | grep -E \"22|3389\"",
        "1  ACCEPT  tcp  --  10.30.8.0/24  0.0.0.0/0  tcp dpt:22",
        "2  DROP    tcp  --  0.0.0.0/0     0.0.0.0/0  tcp dpt:22",
        "3  DROP    tcp  --  0.0.0.0/0     0.0.0.0/0  tcp dpt:3389",
    ]),
    ("g4191", 4924, "综合判定。上述核查结果均符合", [
        "PS C:\\Windows\\system32> quser",
        " 用户名     会话名      状态    空闲时间   登录时间",
        " zhang_wei  rdp-tcp#3   运行中  12         2026/9/12 8:30",
        " li_qiang   console     运行中  1:20       2026/9/12 9:02",
        "PS C:\\Windows\\system32> net accounts",
        "锁定阈值:                                5",
        "锁定持续时间(分):                        30",
        "PS C:\\Windows\\system32> reg query \"HKLM\\SOFTWARE\\Microsoft\\Windows\\CurrentVersion\\Policies\\System\" /v InactivityTimeoutSecs",
        "    InactivityTimeoutSecs    REG_DWORD    0x384",
    ]),
    ("g4193", 4942, "【判定】应用层与操作系统层均需", [
        "[root@kylin-server ~]# grep -r \"loginFail\\|lockout\\|maxAttempts\\|captcha\" /opt/*/conf/*.conf 2>/dev/null",
        "/opt/app/conf/security.yml: login.max-attempts: 5",
        "/opt/app/conf/security.yml: login.lockout-seconds: 1800",
        "/opt/app/conf/security.yml: login.captcha-after: 3",
        "[root@kylin-server ~]# grep -r \"session-timeout\" /opt/app/conf/web.xml",
        "<session-timeout>15</session-timeout>",
    ]),
    ("g4203", 4967, "【判定】应用应尽量采用自主设计", [
        "[root@kylin-server ~]# curl -s -o /dev/null -w \"%{http_code}\\n\" http://localhost/swagger",
        "404",
        "[root@kylin-server ~]# curl -s -o /dev/null -w \"%{http_code}\\n\" http://localhost/actuator",
        "404",
        "[root@kylin-server ~]# find /etc -name \"*.conf\" | xargs grep -l \"sm4\\|gmssl\" 2>/dev/null",
        "/opt/app/conf/protocol.conf",
    ]),
    ("g4213", 4987, "【判定】应用需启用基于数字证书", [
        "[root@kylin-server ~]# grep -r \"clientAuth\\|client_cert\" /opt/*/conf/*.conf 2>/dev/null",
        "/opt/app/conf/tls.conf: client_auth = required",
        "/opt/app/conf/tls.conf: ca_bundle = /etc/app/keys/ca.pem",
        "/opt/app/conf/tls.conf: ukey.pin-verify = on",
    ]),
    ("g4223", 4995, "【判定】Web应用与数据库服务", [
        "[root@kylin-server ~]# grep -E \"listen\" /etc/nginx/nginx.conf /etc/nginx/conf.d/*.conf",
        "nginx.conf:    listen 8843 ssl;",
        "[root@kylin-server ~]# grep -E \"port=\" /opt/tomcat/conf/server.xml | head -2",
        "<Server port=\"8105\" shutdown=\"SHUTDOWN\">",
        "    <Connector port=\"18443\" protocol=\"HTTP/1.1\"/>",
        "[root@kylin-server ~]# sudo ss -tlnp | grep -E \"nginx|java\"",
        "LISTEN 0 511 0.0.0.0:8843  users:((\"nginx\",pid=1902,fd=6))",
        "LISTEN 0 100 *:18443      users:((\"java\",pid=2401,fd=44))",
    ]),
    ("g4233", 5017, "【判定】Web服务器与数据库的管理端口", [
        "[root@kylin-server ~]# sudo ss -tlnp | grep -E ':(80|443|8080|8443|9090)'",
        "LISTEN 0 511 0.0.0.0:443   users:((\"nginx\",pid=1902,fd=6))",
        "LISTEN 0 100 0.0.0.0:8443  users:((\"java\",pid=2401,fd=41))   # 管理后台",
        "mysql> SHOW VARIABLES LIKE 'port';",
        "+---------------+-------+",
        "| Variable_name | Value |",
        "+---------------+-------+",
        "| port          | 3307  |",
        "+---------------+-------+",
    ]),
    ("g4243", 5045, "【判定】应用服务与数据存储", [
        "[root@kylin-server ~]# grep -r \"jdbc\\|Server=\" /opt/*/conf/*.conf 2>/dev/null | grep -v \"127.0.0.1\\|localhost\"",
        "/opt/app/conf/datasource.yml: url: jdbc:mysql://10.30.8.31:3307/appdb?useSSL=true",
        "/opt/app/conf/datasource.yml: host: 10.30.8.31",
    ]),
    ("g4253", 5076, "【判定】系统审计、应用访问日志", [
        "[root@kylin-server ~]# ls -lh /var/log/nginx/access.log",
        "-rw-r----- 1 nginx adm 82M 9月 12 15:20 /var/log/nginx/access.log",
        "[root@kylin-server ~]# grep -rE \"rotate|maxage|180\" /etc/logrotate.d/nginx",
        "/var/log/nginx/*.log { daily rotate 26 maxage 180 }",
        "mysql> SHOW VARIABLES LIKE 'general_log%';",
        "+------------------+----------------------------+",
        "| general_log      | ON                         |",
        "| general_log_file | /var/log/mysql/audit.log   |",
        "+------------------+----------------------------+",
    ]),
    ("g4261", 5083, "核查代码级安全漏洞挖掘过程的规范性", [
        "PS C:\\Windows\\system32> Get-ChildItem -Path D:\\ -Recurse -Include *scan*,*sonar*,*audit* -ErrorAction SilentlyContinue | Select-Object FullName -First 4",
        "",
        "FullName",
        "--------",
        "D:\\code-review\\sonar-report-20260815.html",
        "D:\\code-review\\sonar-report-20260815.pdf",
        "D:\\code-review\\sast-summary-20260815.xlsx",
        "D:\\audit\\vuln-fix-record-20260901.docx",
    ]),
    ("g4263", 5094, "【判定】应用需经过SAST/DAST", [
        "[root@kylin-server ~]# find /opt /home -name \"*scan*\" -o -name \"*报告*\" 2>/dev/null | head -3",
        "/opt/ci/reports/sonar-report-20260901.html",
        "/opt/ci/reports/sast-summary-20260901.pdf",
        "/home/auditor/代码审计记录-20260905.xlsx",
    ]),
    ("g4273", 5117, "【判定】输入格式与长度检查", [
        "[root@kylin-server ~]# grep -rnE \"validate|regex|校验\" /opt/app/src/ 2>/dev/null | head -2",
        "/opt/app/src/web/ValidInterceptor.java:  @Length(max=64)   # 后端独立校验",
        "/opt/app/src/web/ValidInterceptor.java:  Pattern.matches(REGEX_NAME, name)",
    ]),
    ("g4283", 5143, "【判定】应用需在代码层参数化", [
        "[root@kylin-server ~]# grep -rE \"limit_req|limit_conn\" /etc/nginx/ | grep -v \"#\" | head -1",
        "conf.d/app.conf:  limit_req zone=one burst=20 nodelay;",
        "[root@kylin-server ~]# grep -rn \"PreparedStatement\" /opt/app/src/dao/ | head -1",
        "/opt/app/src/dao/UserDao.java:  PreparedStatement ps = conn.prepareStatement(SQL);",
        "[root@kylin-server ~]# grep -c \"XSS\" /var/log/app/waf_audit.log",
        "412",
    ]),
    ("g4293", 5167, "【判定】应用接入统一认证", [
        "[root@kylin-server ~]# systemctl list-units | grep -iE \"iam|cas|sso|oauth\"",
        "cas-server.service   loaded active running   CAS 统一认证服务",
        "[root@kylin-server ~]# grep -r \"sso\\|cas\\|oauth\" /opt/*/conf/*.conf 2>/dev/null | head -2",
        "/opt/app/conf/app.conf: auth.server = https://sso.jy.local/cas",
        "/opt/app/conf/app.conf: auth.mode = cas3",
    ]),
    ("g4303", 5192, "【判定】应用需具备统一管理措施", [
        "[root@kylin-server ~]# sudo ss -tlnp | grep -E \"9090|8443\"",
        "LISTEN 0 100 10.30.8.21:8443 users:((\"java\",pid=2401,fd=41))",
        "[root@kylin-server ~]# sudo iptables -L INPUT -n | grep 8443",
        "ACCEPT  tcp  --  10.30.8.0/24  0.0.0.0/0  tcp dpt:8443",
        "DROP    tcp  --  0.0.0.0/0     0.0.0.0/0  tcp dpt:8443",
    ]),
    ("g4313", 5219, "【判定】最大并发会话数", [
        "[root@kylin-server ~]# grep -r \"maxSession\\|rate\\|并发\" /opt/*/conf/*.conf 2>/dev/null | head -3",
        "/opt/app/conf/server.yml: max-sessions: 500",
        "/opt/app/conf/server.yml: session-rate: 20/s",
        "/opt/app/conf/server.yml: per-user-sessions: 5",
        "[root@kylin-server ~]# grep -r \"limit_conn\\|limit_req\" /etc/nginx/ | grep -v \"#\"",
        "nginx.conf:      limit_req_zone $binary_remote_addr zone=one:10m rate=20r/s;",
        "conf.d/app.conf: limit_conn addr 200;",
    ]),
    ("g4323", 5230, "【判定】应用自身备份功能", [
        "[root@kylin-server ~]# crontab -l | grep mysqldump",
        "0 2 * * * /opt/scripts/backup_mysql.sh",
        "[root@kylin-server ~]# find /opt /root -name \"*restore*\" -o -name \"*recover*\" 2>/dev/null",
        "/opt/scripts/restore_mysql.sh",
        "[root@kylin-server ~]# ls -lh /data/backup/ | tail -2",
        "-rw-r--r-- 1 root root 1.4G 9月 11 02:00 appdb_20260911.sql.gz",
        "-rw-r--r-- 1 root root 1.4G 9月 12 02:00 appdb_20260912.sql.gz",
    ]),
    ("g4333", 5243, "【判定】应用软件需基于国产自主可控", [
        "[root@kylin-server ~]# cat /etc/os-release | head -4",
        "NAME=\"Kylin Linux Advanced Server\"",
        "VERSION=\"V10 (Sword)\"",
        "ID=\"kylin\"",
        "ID_LIKE=\"rhel,fedora,centos\"",
        "[root@kylin-server ~]# rpm -qa | grep -iE \"kylin|dm8|kingbase|nginx\"",
        "kylin-logo-1.0-3.ky10.noarch",
        "nginx-1.21.5-4.ky10.x86_64",
        "[root@kylin-server ~]# rpm -qa | wc -l",
        "238",
    ]),
]


def main():
    os.makedirs(FIGS, exist_ok=True)
    made = 0
    for name, _idx, _prefix, lines in FIGSPEC:
        term(name, lines)
        made += 1
    print("已生成", made, "张图 ->", FIGS)
    if not APPLY:
        print("（图片已生成；加 --apply 插入文档）")
        return 0

    d = Document(PATH)
    paras = d.paragraphs
    ok = fail = skip = 0
    for name, idx, prefix, _lines in FIGSPEC:
        if idx >= len(paras):
            print("!! 索引越界:", name, idx)
            fail += 1
            continue
        p = paras[idx]
        if not p.text.strip().startswith(prefix):
            print("!! 锚文本不符:", name, idx, repr(p.text.strip()[:24]))
            fail += 1
            continue
        # 幂等：后一段已有图则跳过
        nxt = paras[idx + 1] if idx + 1 < len(paras) else None
        if nxt is not None and nxt._p.findall('.//' + qn('a:blip')) and not nxt.text.strip():
            print("已有图，跳过:", name)
            skip += 1
            continue
        src = os.path.join(FIGS, name + ".png")
        w, h = Image.open(src).size
        w_cm = min(14.64, w / DPI * 2.54)
        img = d.add_paragraph()
        img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img.paragraph_format.first_line_indent = None
        img.add_run().add_picture(src, width=Cm(w_cm))
        p._p.addnext(img._p)
        ok += 1
    if fail:
        print("!! %d 个锚点失败，未保存" % fail)
        return 1
    d.save(PATH)
    print("已插入 %d 张（跳过 %d），保存 %s" % (ok, skip, PATH))
    return 0


if __name__ == "__main__":
    sys.exit(main())
