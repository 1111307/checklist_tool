# ⚠️ 已停用（勿再运行）：本脚本会改动目录/分节结构，用户明确要求「不要动目录」，改动已还原。
#    保留仅作审计留痕；如确需重做页码三段式，须先取得用户同意，并在执行前备份。
# -*- coding: utf-8 -*-
"""清理页脚：目录节与正文节统一为「- 页码 -」，去掉未使用的偶数页/首页页脚引用。

（settings.xml 未启用 evenAndOddHeaders，也未启用 titlePg，因此偶数页与首页页脚引用不生效，
  留着只会在文档里留下两套陈旧页脚内容，一并删除。）

用法：python _fix_guide_footer.py [--apply]
"""
import re
import sys

import docx
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

PATH = "配置核查作业指导书_v2.2.docx"
APPLY = "--apply" in sys.argv


def clear_para(p):
    for ch in list(p._p):
        if ch.tag != qn("w:pPr"):
            p._p.remove(ch)


def write_page_footer(sec, size=10.5):
    sec.footer.is_linked_to_previous = False
    f = sec.footer
    p = f.paragraphs[0]
    for extra in list(f.paragraphs)[1:]:
        extra._p.getparent().remove(extra._p)
    clear_para(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run("- ")
    r1.font.size = Pt(size)
    r1.font.name = "Times New Roman"
    r1._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")

    r2 = p.add_run()
    r2.font.size = Pt(size)
    el = r2._element
    b = el.makeelement(qn("w:fldChar"), {}); b.set(qn("w:fldCharType"), "begin")
    i = el.makeelement(qn("w:instrText"), {}); i.set(qn("xml:space"), "preserve"); i.text = "PAGE"
    e = el.makeelement(qn("w:fldChar"), {}); e.set(qn("w:fldCharType"), "end")
    el.append(b); el.append(i); el.append(e)

    r3 = p.add_run(" -")
    r3.font.size = Pt(size)
    r3.font.name = "Times New Roman"
    r3._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    return "".join(p._p.itertext())


def main():
    d = docx.Document(PATH)
    print("节数:", len(d.sections))
    if len(d.sections) < 3:
        print("节数不足 3，先运行 _fix_guide_pagenum.py")
        return 1

    for idx, sec in enumerate(d.sections):
        sectPr = sec._sectPr
        dropped = 0
        for el in list(sectPr.findall(qn("w:footerReference"))):
            t = el.get(qn("w:type"))
            if t in ("even", "first"):
                sectPr.remove(el)
                dropped += 1
        n_ref = len(sectPr.findall(qn("w:footerReference")))
        print("  节%d: 去掉 even/first 页脚引用 %d 个，剩余引用 %d 个" % (idx + 1, dropped, n_ref))

    if not APPLY:
        print("\n（干跑，未写文件；加 --apply 执行）")
        return 0

    if len(d.sections) >= 3:
        print("  节1 封面：保持无页脚引用")
        print("  节2 页脚 ->", write_page_footer(d.sections[1]))
        print("  节3 页脚 ->", write_page_footer(d.sections[2]))

    d.save(PATH)
    print("已写入", PATH)

    d2 = docx.Document(PATH)
    print("\n=== 校验 ===")
    for i, s in enumerate(d2.sections):
        sx = re.sub(r"\s+", " ", s._sectPr.xml)
        print("  节%d: pgNumType=%s 页脚引用=%d 页脚文本=%r" % (
            i + 1,
            re.findall(r'<w:pgNumType[^>]*/>', sx),
            len(re.findall("footerReference", sx)),
            "".join(s.footer._element.itertext()).strip()[:24]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
