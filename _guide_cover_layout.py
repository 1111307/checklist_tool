# -*- coding: utf-8 -*-
r"""指导书封面重排为用户确认的版式：标题在上 → 中部校徽+校名 → 底部说明行+日期。

做法：搬移原有段落元素（保留格式），仅重排顺序并调整空行数量；校徽 3.5cm → 4.0cm。
"""
import docx
from docx.oxml.ns import qn
from docx.shared import Cm, Emu

PATH = "配置核查作业指导书_v2.0.0.docx"

d = docx.Document(PATH)
ps = d.paragraphs
# 定位封面元素
t1 = next(p for p in ps[:20] if p.text.strip() == "配置核查")
t2 = next(p for p in ps[:20] if p.text.strip() == "作业指导书")
img = next(p for p in ps[:20] if p._p.findall('.//' + qn('w:drawing')))
school = next(p for p in ps[:24] if p.text.strip() == "西北工业大学")
note = next(p for p in ps[:30] if p.text.startswith("适合系统/数据库/中间件"))
date = next(p for p in ps[:30] if p.text.strip().startswith("2026"))
sect_p = next(p for p in ps[:30] if p._p.find('.//' + qn('w:sectPr')) is not None)

# 校徽放大到 4.0cm（方形，等比）
for ext in img._p.findall('.//' + qn('wp:extent')):
    ext.set('cx', str(Emu(Cm(4.0)).emu if hasattr(Emu, 'emu') else int(Cm(4.0))))
    ext.set('cy', str(int(Cm(4.0))))
for ext in img._p.findall('.//' + qn('a:ext')):
    ext.set('cx', str(int(Cm(4.0))))
    ext.set('cy', str(int(Cm(4.0))))

# 删除封面区旧空段（标题/校徽/校名/说明/日期之外的空段）
keep = {id(x._p) for x in (t1, t2, img, school, note, date, sect_p)}
for p in ps[:30]:
    if id(p._p) in keep:
        continue
    if p.text.strip() == "" and p._p.findall('.//' + qn('w:drawing')) == [] and p._p.find('.//' + qn('w:sectPr')) is None:
        p._p.getparent().remove(p._p)

# 按新版式重新插入（元素搬移，格式不变）
anchor = sect_p._p
def new_empty():
    p = d.add_paragraph()
    anchor.addprevious(p._p)
    return p

for _ in range(5):
    new_empty()
anchor.addprevious(t1._p)
anchor.addprevious(t2._p)
for _ in range(8):
    new_empty()
anchor.addprevious(img._p)
anchor.addprevious(school._p)
for _ in range(10):
    new_empty()
anchor.addprevious(note._p)
new_empty()
anchor.addprevious(date._p)

d.save(PATH)

chk = docx.Document(PATH)
print("封面结构（重排后）：")
for i, p in enumerate(chk.paragraphs[:30]):
    imgs = len(p._p.findall('.//' + qn('w:drawing')))
    sect = p._p.find('.//' + qn('w:sectPr')) is not None
    if p.text.strip() or imgs or sect:
        print("  %2d %-30s 图=%d%s" % (i, repr(p.text[:28]), imgs, " 分节" if sect else ""))
