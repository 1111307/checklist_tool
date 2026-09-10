# -*- coding: utf-8 -*-
"""修正核查表两处合并缺陷（不改任何数据）：

1. A4:A29 → A4:A30：补 1.27 行（R30）的章级标签，原来 1.27 悬在「系统安全」合并区之外，
   A 列看不到所属章节
2. N2:N3 合并：标注列表头「工具自动验证」原来只在第 2 行，与 A2:A3/B2:B3 的两行表头不齐
"""
import openpyxl

XLSX = '配置核查表_v2.0.0_标注自动验证.xlsx'
wb = openpyxl.load_workbook(XLSX)
ws = wb['Sheet1']
ranges = [str(m) for m in ws.merged_cells.ranges]

# 1) 系统安全章标签覆盖到 1.27 行
if 'A4:A29' in ranges:
    ws.unmerge_cells('A4:A29')
    assert ws.cell(4, 1).value == '系统安全', ws.cell(4, 1).value
    assert not str(ws.cell(30, 1).value or '').strip(), repr(ws.cell(30, 1).value)
    ws.merge_cells('A4:A30')
    print('已修：A4:A29 → A4:A30（1.27 行归入系统安全）')

# 2) 标注列表头两行合并
if 'N2:N3' not in ranges:
    assert ws.cell(2, 14).value == '工具自动验证', ws.cell(2, 14).value
    assert not str(ws.cell(3, 14).value or '').strip()
    ws.merge_cells('N2:N3')
    print('已修：N2:N3 合并（标注列表头与两行表头对齐）')

wb.save(XLSX)

# 回读校验
wb2 = openpyxl.load_workbook(XLSX); ws2 = wb2['Sheet1']
mg = [str(m) for m in ws2.merged_cells.ranges]
assert 'A4:A30' in mg and 'N2:N3' in mg
assert ws2.cell(4, 1).value == '系统安全' and ws2.cell(2, 14).value == '工具自动验证'
ticks = sum(1 for r in range(4, 140) for c in range(3, 14) if ws2.cell(r, c).value == '√')
print('回读通过：合并', len(mg), '处，勾选总数', ticks)
