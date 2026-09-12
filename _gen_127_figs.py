# -*- coding: utf-8 -*-
r"""1.27.3 / 1.27.4 配图生成并插入指导书。

四张图（_tmp_figs/）：
  gen_nginx.png    Nginx：终端查版本 + 商业版权证文件（Linux 风格）
  gen_tomcat.png   Tomcat：catalina.bat version 版本信息（Windows 风格）
  gen_office.png   Office：slmgr /dli 授权对话框（Windows 弹窗样式）
  gen_wps.png      WPS：帐号页面（窗口示意）

用法：
  python _gen_127_figs.py            # 只生成图片
  python _gen_127_figs.py --apply    # 生成并插入 docx（先自行备份）
"""
import os
import subprocess
import sys

from PIL import Image, ImageChops
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm

HERE = os.path.dirname(os.path.abspath(__file__))
FIGS = os.path.join(HERE, "_tmp_figs")
PATH = os.path.join(HERE, "配置核查作业指导书_v2.2.docx")
CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
APPLY = "--apply" in sys.argv
DPI = 110.0
MONO = 'font:15px/1.9 "Courier New","Consolas",monospace;color:#000;'


def shoot(html, png):
    html_path = os.path.join(FIGS, "_t.html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-sandbox",
                    "--hide-scrollbars", "--force-device-scale-factor=1",
                    "--virtual-time-budget=3000", "--window-size=1200,900",
                    "--screenshot=" + png, "file:///" + html_path.replace("\\", "/")],
                   capture_output=True, text=True)
    os.remove(html_path)


def crop(png, pad=8):
    im = Image.open(png).convert("RGB")
    bg = Image.new("RGB", im.size, (255, 255, 255))
    bbox = ImageChops.difference(im, bg).getbbox()
    if bbox:
        im.crop((max(bbox[0] - pad, 0), max(bbox[1] - pad, 0),
                 min(bbox[2] + pad, im.width), min(bbox[3] + pad, im.height))).save(png)
    return Image.open(png).size


def term(name, lines):
    body = "\n".join(lines).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    html = ('<!DOCTYPE html><html><head><meta charset="UTF-8"><style>'
            'body{margin:0;background:#fff;}pre{%s margin:0;padding:6px 10px;white-space:pre;}'
            '</style></head><body><pre>%s</pre></body></html>' % (MONO, body))
    png = os.path.join(FIGS, name + ".png")
    shoot(html, png)
    print(name, crop(png))


DIALOG_CSS = """
body{margin:0;background:#fff;font-family:"Microsoft YaHei",sans-serif;}
.dlg{width:520px;border:1px solid #7a7a7a;border-radius:4px;box-shadow:2px 3px 8px rgba(0,0,0,.25);}
.tbar{background:linear-gradient(#f5f6f7,#dfe1e3);padding:5px 10px;font:13px "Microsoft YaHei";color:#333;border-bottom:1px solid #bbb;}
.dbody{padding:14px 18px;font:15px/1.9 "Courier New",monospace;color:#000;white-space:pre;}
.btnrow{padding:0 14px 12px;text-align:right;}
.btn{display:inline-block;border:1px solid #999;border-radius:3px;background:#f0f0f0;padding:3px 26px;font:13px "Microsoft YaHei";color:#333;}
"""


def dialog(name, title, body_lines):
    body = "\n".join(body_lines).replace("&", "&amp;").replace("<", "&lt;")
    html = ('<!DOCTYPE html><html><head><meta charset="UTF-8"><style>%s</style></head>'
            '<body><div class="dlg"><div class="tbar">%s</div>'
            '<div class="dbody">%s</div>'
            '<div class="btnrow"><span class="btn">确定</span></div></div></body></html>'
            % (DIALOG_CSS, title, body))
    png = os.path.join(FIGS, name + ".png")
    shoot(html, png)
    print(name, crop(png))


WPS_CSS = """
body{margin:0;background:#fff;font-family:"Microsoft YaHei",sans-serif;}
.win{width:640px;border:1px solid #8a8a8a;border-radius:4px;overflow:hidden;box-shadow:2px 3px 8px rgba(0,0,0,.2);}
.tbar{background:#e8eaed;padding:6px 12px;font:13px "Microsoft YaHei";color:#333;border-bottom:1px solid #c5c7ca;}
.flex{display:flex;}
.nav{width:130px;background:#f4f5f7;border-right:1px solid #d9dbde;padding:10px 0;}
.nav div{padding:8px 22px;font:14px "Microsoft YaHei";color:#444;}
.nav .on{background:#e2e6ea;color:#111;font-weight:bold;}
.main{padding:26px 30px;flex:1;}
.ava{width:56px;height:56px;border-radius:50%;background:#c9cdd2;color:#fff;text-align:center;line-height:56px;font-size:24px;}
.uname{font:16px "Microsoft YaHei";color:#222;margin-top:10px;}
.kv{font:13px "Microsoft YaHei";color:#555;margin-top:14px;}
.kv b{color:#111;font-weight:normal;border:1px solid #b8bcc0;border-radius:3px;padding:1px 8px;font-size:12px;}
"""


def wps(name):
    html = ('<!DOCTYPE html><html><head><meta charset="UTF-8"><style>%s</style></head><body>'
            '<div class="win"><div class="tbar">WPS Office</div><div class="flex">'
            '<div class="nav"><div>首页</div><div>文档</div><div>表格</div><div>演示</div>'
            '<div class="on">帐号</div></div>'
            '<div class="main"><div class="ava">张</div>'
            '<div class="uname">zhang_wei@wps.cn</div>'
            '<div class="kv">会员状态：<b>WPS 超级会员（个人版）</b></div>'
            '<div class="kv">有效期至：2027-03-15</div>'
            '<div class="kv">授权凭证：企业采购序列号（已注册）</div>'
            '</div></div></div></body></html>' % WPS_CSS)
    png = os.path.join(FIGS, name + ".png")
    shoot(html, png)
    print(name, crop(png, pad=0))


# ---------------- 生成 ----------------
term("gen_nginx", [
    "[root@kylin-server ~]# nginx -v",
    "nginx version: nginx/1.26.2",
    '[root@kylin-server ~]# ls /etc/nginx/nginx-repo.*',
    "ls: cannot access '/etc/nginx/nginx-repo.*': No such file or directory",
    "[root@kylin-server ~]#",
])

term("gen_tomcat", [
    r"C:\Tomcat\bin> catalina.bat version",
    'Using CATALINA_BASE:   "C:\\Tomcat"',
    'Using CATALINA_HOME:   "C:\\Tomcat"',
    'Server version:  Apache Tomcat/9.0.85',
    'Server built:    Jan 10 2024 12:00:00 UTC',
    'Server number:   9.0.85.0',
    'OS Name:         Windows Server 2019',
    'JVM Version:     1.8.0_392',
])

dialog("gen_office", "Windows Script Host", [
    "软件授权服务管理",
    "",
    "名称: Microsoft Office Professional Plus 2016",
    "描述: Microsoft Office Professional Plus 2016",
    "部分产品密钥: XVK2C",
    "许可证状态: 已授权",
])

wps("gen_wps")


# ---------------- 插入 ----------------
def insert():
    d = Document(PATH)
    paras = d.paragraphs

    def find_para(prefix):
        for p in paras:
            if p.text.strip().startswith(prefix):
                return p
        return None

    jobs = [
        ("核查Nginx是否具备完备的售后技术支持与服务", "gen_nginx.png"),
        ("核查Tomcat是否具备完备的售后技术支持与服务", "gen_tomcat.png"),
        ("任务栏搜索框输入“cmd”或“命令提示符”", "gen_office.png"),
        ("打开WPS Office，点击左上角", "gen_wps.png"),
    ]
    for prefix, fig in jobs:
        src = os.path.join(FIGS, fig)
        w, h = Image.open(src).size
        w_cm = min(14.64, w / DPI * 2.54)
        anchor = find_para(prefix)
        if anchor is None:
            print("!! 未找到:", prefix[:20])
            continue
        # 已插入过则跳过（下一段是图片段落即视为已插入）
        nxt = None
        for p in paras:
            if p._p is anchor._p:
                idx = paras.index(p)
                nxt = paras[idx + 1] if idx + 1 < len(paras) else None
                break
        if nxt is not None and nxt._p.findall('.//' + qn('a:blip')):
            print("跳过（已有图）:", prefix[:20])
            continue
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = None
        p.add_run().add_picture(src, width=Cm(w_cm))
        anchor._p.addnext(p._p)
        print("已插入 %s 于 %r 后（%.2fcm）" % (fig, prefix[:14], w_cm))

    if APPLY:
        d.save(PATH)
        print("已保存", PATH)
    else:
        print("（未保存；加 --apply 写入）")


if __name__ == "__main__":
    if APPLY:
        insert()
