# -*- coding: utf-8 -*-
"""收尾修正两处：

1) 26 个「平台子标题」的字体经主题引用解析成了宋体（与同类标题的黑体不一致）：
   给这些标题的 run 与段落标记补显式字体（黑体 / Times New Roman），并清掉主题字体引用；
2) 「Redis没有存储过程。」由标题改回正文后仍带加粗：去掉加粗。

用法：python _fix_guide_fonts2.py [--apply]
"""
import re
import sys

import docx
from docx.oxml.ns import qn

PATH = "配置核查作业指导书_v2.2.docx"
APPLY = "--apply" in sys.argv
SIZES = {"1": 16.0, "2": 14.0, "3": 12.0, "4": 10.5, "5": 10.5}
THEME_ATTRS = ("w:asciiTheme", "w:eastAsiaTheme", "w:hAnsiTheme", "w:cstheme")


def fix_rpr(rPr, east, size, latin="Times New Roman"):
    """写入显式字体/字号，并清掉主题字体引用。"""
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rf)
    for a in THEME_ATTRS:
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    rf.set(qn("w:ascii"), latin)
    rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:eastAsia"), east)
    sz = rPr.find(qn("w:sz"))
    if sz is None:
        sz = rPr.makeelement(qn("w:sz"), {})
        rf.addnext(sz)
    sz.set(qn("w:val"), str(int(size * 2)))
    szcs = rPr.find(qn("w:szCs"))
    if szcs is not None:
        szcs.set(qn("w:val"), str(int(size * 2)))


def main():
    d = docx.Document(PATH)
    heads_fixed = []
    redis_fixed = False

    for p in d.paragraphs:
        st = p.style.name
        is_head = st.startswith("Heading")
        txt = p.text.strip()
        if not txt:
            continue

        # ---- 1) 平台子标题：缺显式中文字体 ----
        if is_head:
            lvl = st.split()[-1]
            size = SIZES.get(lvl, 12.0)
            miss = False
            for r in p.runs:
                rPr = r._element.rPr
                rf = rPr.find(qn("w:rFonts")) if rPr is not None else None
                if rf is None or rf.get(qn("w:eastAsia")) is None:
                    miss = True
                    break
            if miss:
                heads_fixed.append(txt[:42])
                if APPLY:
                    for r in p.runs:
                        fix_rpr(r._element.get_or_add_rPr(), "黑体", size)
                    pPr = p._p.find(qn("w:pPr"))
                    if pPr is not None:
                        mark = pPr.find(qn("w:rPr"))
                        if mark is not None:
                            fix_rpr(mark, "黑体", size)

        # ---- 2) Redis 那句：去掉加粗 ----
        if txt.startswith("Redis没有存储过程"):
            redis_fixed = True
            if APPLY:
                for r in p.runs:
                    rPr = r._element.get_or_add_rPr()
                    for b in rPr.findall(qn("w:b")):
                        rPr.remove(b)
                    for b in rPr.findall(qn("w:bCs")):
                        rPr.remove(b)
                rPr = p._p.get_or_add_pPr().find(qn("w:rPr"))
                if rPr is not None:
                    for b in rPr.findall(qn("w:b")) + rPr.findall(qn("w:bCs")):
                        rPr.remove(b)

    print("=== 1. 补显式中文字体的标题（%d 个）===" % len(heads_fixed))
    for t in heads_fixed[:10]:
        print("   ", t)
    if len(heads_fixed) > 10:
        print("    ...")
    print("=== 2. 去掉「Redis没有存储过程。」的加粗:", "已定位" if redis_fixed else "未找到")

    if not APPLY:
        print("\n（干跑，未写文件；加 --apply 执行）")
        return 0

    d.save(PATH)
    print("\n已写入", PATH)

    # 复核
    d2 = docx.Document(PATH)
    miss = 0
    for p in d2.paragraphs:
        if not p.style.name.startswith("Heading") or not p.text.strip():
            continue
        for r in p.runs:
            rPr = r._element.rPr
            rf = rPr.find(qn("w:rFonts")) if rPr is not None else None
            if rf is None or rf.get(qn("w:eastAsia")) is None:
                miss += 1
                break
    print("复核：仍缺显式中文字体的标题数 =", miss)
    for p in d2.paragraphs:
        if p.text.strip().startswith("Redis没有存储过程"):
            print("复核：Redis 段加粗 =", any(r.bold for r in p.runs))
            break
    return 0


if __name__ == "__main__":
    sys.exit(main())
