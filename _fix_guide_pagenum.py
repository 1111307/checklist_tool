# ⚠️ 已停用（勿再运行）：本脚本会改动目录/分节结构，用户明确要求「不要动目录」，改动已还原。
#    保留仅作审计留痕；如确需重做页码三段式，须先取得用户同意，并在执行前备份。
# -*- coding: utf-8 -*-
"""指导书页码改为三段式：封面无页码 → 目录小写罗马数字 → 正文阿拉伯数字从 1 起。

做法：
  1. 在封面最后一段挂 sectPr（下一页分节）→ 封面独立成节，且不带页脚引用（无页码）
  2. 在目录内容控件之后插入一个带 sectPr 的空段落 → 目录独立成节，页码 fmt=lowerRoman start=1
  3. 正文沿用文档末尾的 sectPr，页码 fmt=decimal start=1
  4. 给目录节与正文节各自建页脚，写入「- PAGE -」域

用法：
  python _fix_guide_pagenum.py            # 干跑（只打印将要做的改动）
  python _fix_guide_pagenum.py --apply
"""
import copy
import re
import shutil
import sys

import docx
from docx.oxml.ns import qn

PATH = "配置核查作业指导书_v2.2.docx"
BAK = "../_guide_backup/配置核查作业指导书_v2.2.beforepagenum.docx"
APPLY = "--apply" in sys.argv
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def rid_of(el):
    return el.get(qn("r:id"))


def main():
    d = docx.Document(PATH)
    body = d.element.body
    sect = body.find(qn("w:sectPr"))
    if sect is None:
        print("未找到文档级 sectPr")
        return 1

    # 定位封面末段与目录 sdt
    kids = list(body)
    sdt = next((el for el in kids if el.tag == W + "sdt"), None)
    if sdt is None:
        print("未找到目录内容控件")
        return 2
    idx = kids.index(sdt)
    prev_p = None
    for el in reversed(kids[:idx]):
        if el.tag == W + "p":
            prev_p = el
            break
    if prev_p is None:
        print("未找到封面末段")
        return 3

    print("目录 sdt 位置:", idx, "| 封面末段文本:", repr("".join(prev_p.itertext())[:20]))

    # 已有的分节数
    n_sec = len(d.sections)
    print("当前节数:", n_sec)
    if n_sec > 1:
        print("已有多节，本脚本只处理单节情形，退出。")
        return 4

    # ---- 构造两个分节属性（复制文档级 sectPr 作为模板）----
    s1 = copy.deepcopy(sect)          # 封面：去掉全部页脚引用 → 无页码
    for tag in ("footerReference", "headerReference"):
        for el in s1.findall(qn("w:" + tag)):
            s1.remove(el)
    # 封面节不需要 pgNumType
    for el in s1.findall(qn("w:pgNumType")):
        s1.remove(el)
    s1.set(qn("w:type"), "nextPage") if s1.get(qn("w:type")) is None else None

    s2 = copy.deepcopy(sect)          # 目录：小写罗马数字，从 1 起
    for el in s2.findall(qn("w:pgNumType")):
        s2.remove(el)
    pn = s2.makeelement(qn("w:pgNumType"), {})
    pn.set(qn("w:fmt"), "lowerRoman")
    pn.set(qn("w:start"), "1")
    s2.append(pn)
    for tag in ("footerReference", "headerReference"):   # 页脚后面用 python-docx 重建
        for el in s2.findall(qn("w:" + tag)):
            s2.remove(el)

    s3 = sect                          # 正文：阿拉伯，从 1 起
    for el in s3.findall(qn("w:pgNumType")):
        s3.remove(el)
    pn3 = s3.makeelement(qn("w:pgNumType"), {})
    pn3.set(qn("w:fmt"), "decimal")
    pn3.set(qn("w:start"), "1")
    s3.append(pn3)

    print("\n=== 计划 ===")
    print("  节1 封面：挂 sectPr 在封面末段，无页脚 → 不显示页码")
    print("  节2 目录：目录内容控件后插入空段落并挂 sectPr，页码 lowerRoman 从 1 起")
    print("  节3 正文：文档末尾 sectPr，页码 decimal 从 1 起")

    if not APPLY:
        print("\n（干跑，未写文件；加 --apply 执行）")
        return 0

    shutil.copy(PATH, BAK)
    print("\n已备份 ->", BAK)

    # 1) 封面末段挂 sectPr
    pPr = prev_p.find(qn("w:pPr"))
    if pPr is None:
        pPr = prev_p.makeelement(qn("w:pPr"), {})
        prev_p.insert(0, pPr)
    pPr.append(s1)

    # 2) 目录 sdt 之后插入带 sectPr 的空段落
    newp = sdt.makeelement(qn("w:p"), {})
    npPr = sdt.makeelement(qn("w:pPr"), {})
    npPr.append(s2)
    newp.append(npPr)
    sdt.addnext(newp)

    d.save(PATH)
    print("已写入", PATH)

    # 3) 用 python-docx 重新打开，给节2/节3 建页脚
    d2 = docx.Document(PATH)
    print("写入后节数:", len(d2.sections))

    def set_footer(sec, text_around="-"):
        sec.footer.is_linked_to_previous = False
        f = sec.footer
        # 清空后写入「- PAGE -」
        for p in list(f.paragraphs)[1:]:
            p._p.getparent().remove(p._p)
        p = f.paragraphs[0]
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text_around + " ")
        r2 = p.add_run()
        fld_b = r2._element.makeelement(qn("w:fldChar"), {}); fld_b.set(qn("w:fldCharType"), "begin")
        instr = r2._element.makeelement(qn("w:instrText"), {}); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
        fld_e = r2._element.makeelement(qn("w:fldChar"), {}); fld_e.set(qn("w:fldCharType"), "end")
        r2._element.append(fld_b); r2._element.append(instr); r2._element.append(fld_e)
        p.add_run(" " + text_around)
        return p.text

    if len(d2.sections) >= 3:
        print("  节1 页脚保持为空（无引用）")
        set_footer(d2.sections[1])
        set_footer(d2.sections[2])
        # 封面节确保没有页脚引用
        s1sect = d2.sections[0]._sectPr
        for el in s1sect.findall(qn("w:footerReference")):
            s1sect.remove(el)
        d2.save(PATH)
        print("页脚已写入并保存")
    else:
        print("节数不为 3，未写页脚，请检查")

    # 校验
    d3 = docx.Document(PATH)
    print("\n=== 校验 ===")
    for i, s in enumerate(d3.sections):
        sx = re.sub(r"\s+", " ", s._sectPr.xml)
        fmt = re.findall(r'<w:pgNumType[^>]*>', sx)
        fref = len(re.findall(r"footerReference", sx))
        ftxt = "".join(s.footer._element.itertext()).strip()
        print("  节%d: pgNumType=%s 页脚引用=%d 页脚文本=%r" % (i + 1, fmt, fref, ftxt[:20]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
