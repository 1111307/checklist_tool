# -*- coding: utf-8 -*-
"""指导书格式批量修正（不动正文文字内容，只改样式/字体/字号/编号空格）。

修正项：
  A 1.8.4 后的「Redis没有存储过程。」由 Heading 4 改回正文样式
  B 2.16.3 由 Heading 4 升为 Heading 3（与 2.16.1/2.16.2 同级）
  C 补全缺失的字体设置：标题缺中文字体→黑体，正文缺→宋体，西文缺→Times New Roman
  D 标题编号与标题文字之间补空格（统一为「1.1 标题」形式）
  E 正文字号 15pt → 14pt
  F 代码块（HTML Preformatted）字号 6.5pt → 10.5pt、行距改单倍
  G 网页样式 Normal (Web) → 正文样式 Normal

用法：
  python _fix_guide_format.py            # 干跑预览（不写文件）
  python _fix_guide_format.py --apply    # 实际写入
"""
import re
import shutil
import sys

import docx
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

PATH = "配置核查作业指导书_v2.2.docx"
BAK = "../_guide_backup/配置核查作业指导书_v2.2.beforeformat.docx"
APPLY = "--apply" in sys.argv

BODY_SIZE = 14.0
CODE_SIZE = 10.5


def set_run_font(r, east, size=None, latin="Times New Roman"):
    rPr = r._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rf)
    if rf.get(qn("w:ascii")) is None:
        rf.set(qn("w:ascii"), latin)
    if rf.get(qn("w:hAnsi")) is None:
        rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:eastAsia"), east)
    if size is not None:
        r.font.size = Pt(size)


def insert_space_after_number(p):
    """在标题编号与文字之间插入一个半角空格（按字符位置落到具体 run）。"""
    full = p.text
    m = re.match(r"^(\d+(?:\.\d+)+|\d+)(\S)", full)
    if not m:
        return False
    pos = len(m.group(1))          # 空格插入位置
    acc = 0
    for r in p.runs:
        nxt = acc + len(r.text)
        if acc <= pos <= nxt:
            off = pos - acc
            r.text = r.text[:off] + " " + r.text[off:]
            return True
        acc = nxt
    return False


def main():
    if APPLY:
        shutil.copy(PATH, BAK)
        print("已备份 ->", BAK)

    d = docx.Document(PATH)
    todo = []

    # ---------- A / B：两处样式错误 ----------
    target_redis = None
    target_2163 = None
    prev_normal = None
    for p in d.paragraphs:
        t = p.text.strip()
        if p.style.name == "Heading 4" and t.startswith("Redis没有存储过程"):
            target_redis = p
        if p.style.name == "Heading 4" and t.startswith("2.16.3"):
            target_2163 = p
        if p.style.name == "Normal" and t and prev_normal is None:
            prev_normal = p
    if target_redis is not None:
        todo.append(("A 改正文样式", target_redis.text.strip()[:40]))
        if APPLY:
            target_redis.style = d.styles["Normal"]
            pf = target_redis.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            ind = target_redis._p.get_or_add_pPr().get_or_add_ind()
            ind.set(qn("w:firstLineChars"), "200")
            ind.set(qn("w:firstLine"), str(int(BODY_SIZE * 2 * 20)))
            for r in target_redis.runs:
                set_run_font(r, "宋体", BODY_SIZE)
    if target_2163 is not None:
        todo.append(("B 升为三级标题", target_2163.text.strip()[:40]))
        if APPLY:
            target_2163.style = d.styles["Heading 3"]
            for r in target_2163.runs:
                set_run_font(r, "黑体", 12.0)
                r.bold = True

    # ---------- C：补字体 ----------
    miss_h = miss_b = 0
    for p in d.paragraphs:
        is_head = p.style.name.startswith("Heading")
        for r in p.runs:
            if not r.text.strip():
                continue
            rPr = r._element.rPr
            rf = rPr.rFonts if rPr is not None else None
            if rf is None or rf.get(qn("w:eastAsia")) is None:
                if is_head:
                    miss_h += 1
                    if APPLY:
                        set_run_font(r, "黑体")
                else:
                    miss_b += 1
                    if APPLY:
                        set_run_font(r, "宋体")
    todo.append(("C 补中文字体（标题 run / 正文 run）", "%d / %d" % (miss_h, miss_b)))

    # ---------- D：编号后补空格 ----------
    fixed = []
    for p in d.paragraphs:
        if not p.style.name.startswith("Heading") or not p.text.strip():
            continue
        t = p.text.strip()
        m = re.match(r"^(\d+(?:\.\d+)+|\d+)(.*)$", t)
        if m and m.group(2)[:1] not in (" ", "\u3000"):
            fixed.append(t[:40])
            if APPLY:
                insert_space_after_number(p)
    todo.append(("D 编号后补空格", "%d 处，例：%s" % (len(fixed), fixed[0] if fixed else "")))

    # ---------- E：正文字号 15pt ----------
    n15 = 0
    for p in d.paragraphs:
        if p.style.name.startswith("Heading"):
            continue
        for r in p.runs:
            if r.font.size and abs(r.font.size.pt - 15.0) < 0.01:
                n15 += 1
                if APPLY:
                    r.font.size = Pt(BODY_SIZE)
    todo.append(("E 正文字号 15pt→14pt", "%d 处" % n15))

    # ---------- F：代码块字号 ----------
    npre = 0
    for p in d.paragraphs:
        if p.style.name == "HTML Preformatted":
            npre += 1
            if APPLY:
                for r in p.runs:
                    r.font.size = Pt(CODE_SIZE)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(2)
    todo.append(("F 代码块字号 6.5pt→%.1fpt" % CODE_SIZE, "%d 段" % npre))

    # ---------- G：网页样式归正 ----------
    nweb = 0
    for p in d.paragraphs:
        if p.style.name == "Normal (Web)":
            nweb += 1
            if APPLY:
                p.style = d.styles["Normal"]
    todo.append(("G Normal (Web)→Normal", "%d 段" % nweb))

    print("\n=== 待执行 ===")
    for k, v in todo:
        print("  %-34s %s" % (k, v))

    if APPLY:
        d.save(PATH)
        print("\n已写入", PATH)
    else:
        print("\n（干跑，未写文件；加 --apply 执行）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
