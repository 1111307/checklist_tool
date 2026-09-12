# -*- coding: utf-8 -*-
r"""1.24.2.3 达梦 / 1.24.2.4 Redis 两处整理：

A. 1.24.2.3 达梦：在 Windows 版本的方法段之后补一张 disql 登录示例图（用户要求"windows版本最后生成一张图"）
B. 1.24.2.4 Redis：
   B1. 两个平台伪标题补首行缩进（前面各小节的「…操作系统下的核查方法」都有 firstLine=560，
       Redis 这两处只有 firstLineChars，视觉上没对齐）
   B2. 调整内容归属与顺序，与前面小节（Windows 组在前、麒麟组在后）一致：
       原结构为 [3 段无归属方法][麒麟伪标题][4 段][Windows 内容][Windows 伪标题（下面空）]
       改为 [Windows 伪标题][Windows 内容][图][麒麟伪标题][3 段无归属方法归入麒麟][4 段][图]
   B3. 两个平台组各补一张示例图

用法：
  python _fix_124_sections.py            # 干跑
  python _fix_124_sections.py --apply
"""
import os
import shutil
import sys

import docx
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

PATH = "配置核查作业指导书_v2.2.docx"
BAK = "../_guide_backup/配置核查作业指导书_v2.2.before124sections.docx"
HERE = os.path.dirname(os.path.abspath(__file__))
IMG_WIN_DM = os.path.join(HERE, "_tmp_imgs", "gen_dm_win.png")
IMG_WIN_REDIS = os.path.join(HERE, "_tmp_imgs", "gen_redis_win.png")
IMG_KYLIN_REDIS = os.path.join(HERE, "_tmp_imgs", "gen_redis_kylin.png")
APPLY = "--apply" in sys.argv
IMG_CM = 14.64


def new_img_para(doc, path):
    """新建居中的图片段落（追加到文末，随后用元素引用搬到目标位置）。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.add_run().add_picture(path, width=Cm(IMG_CM))
    return p._p


def set_first_line_indent(p, twips=560):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = pPr.makeelement(qn("w:ind"), {})
        pPr.append(ind)
    ind.set(qn("w:firstLine"), str(twips))
    # 与前面各小节保持一致：firstLineChars 同时保留
    if ind.get(qn("w:firstLineChars")) is None:
        ind.set(qn("w:firstLineChars"), "200")


def main():
    for f in (IMG_WIN_DM, IMG_WIN_REDIS, IMG_KYLIN_REDIS):
        if not os.path.exists(f):
            print("缺少图片:", f)
            return 2

    d = docx.Document(PATH)
    paras = d.paragraphs

    def find(pred, start=0):
        for i in range(start, len(paras)):
            if pred(paras[i].text.strip()):
                return i
        return None

    # 定位关键段落
    i_dm_win_method = find(lambda t: t.startswith("使用win+R组合键打开运行框，输入cmd，在命令行进入达梦安装目录"),
                           3200)
    i_redis_h = find(lambda t: t.startswith("1.24.2.4 Redis"))
    assert i_redis_h, "未找到 1.24.2.4 Redis"
    i_stat = i_redis_h + 1                     # 判定标准
    i_a = find(lambda t: t.startswith('输入“grep -E "logfile|loglevel"'), i_stat)
    i_b = find(lambda t: t.startswith("输入“CONFIG GET dir”"), i_stat)
    i_c = find(lambda t: t.startswith("在桌面鼠标右键-打开终端，命令行输入“redis-cli -a 密码”"), i_stat)
    i_kt = find(lambda t: t == "中标麒麟、银河麒麟操作系统下的核查方法", i_stat)
    i_k1 = find(lambda t: t.startswith("核查日志保留策略"), i_stat)
    i_k4 = find(lambda t: t.startswith("核查日志开启情况"), i_stat)
    i_w1 = find(lambda t: t.startswith("打开Redis安装目录"), i_stat)
    i_wt = find(lambda t: t == "Windows操作系统下的核查方法", i_stat)
    print("定位: 达梦Win方法=%s Redis判定=%s A=%s B=%s C=%s 麒麟标题=%s K1=%s K4=%s Win内容=%s Win标题=%s"
          % (i_dm_win_method, i_stat, i_a, i_b, i_c, i_kt, i_k1, i_k4, i_w1, i_wt))
    assert None not in (i_dm_win_method, i_a, i_b, i_c, i_kt, i_k1, i_k4, i_w1, i_wt)

    order_ok = i_a < i_b < i_c < i_kt < i_k1 <= i_k4 < i_w1 < i_wt
    print("原顺序符合预期:", order_ok)
    if not order_ok:
        print("!! 段序与预期不符，请人工确认后再执行")
        return 3

    print("\n=== 计划 ===")
    print("A. 达梦：在『…进入达梦安装目录的bin文件夹…』段后插入 disql 登录示例图")
    print("B1. Redis：为『中标麒麟、银河麒麟操作系统下的核查方法』与『Windows操作系统下的核查方法』补首行缩进(560 twips)")
    print("B2. Redis 段序调整为：")
    print("    判定标准 → [Windows伪标题] → [Windows内容] → [Win图] → [麒麟伪标题] → A/B/C 三段 → K1~K4 → [麒麟图]")
    print("B3. 插入 2 张示例图（Windows / 麒麟）")

    if not APPLY:
        print("\n（干跑，未写文件；加 --apply 执行）")
        return 0

    shutil.copy(PATH, BAK)
    print("\n已备份 ->", BAK)

    # ---- A. 达梦补图 ----
    img_dm = new_img_para(d, IMG_WIN_DM)
    paras[i_dm_win_method]._p.addnext(img_dm)

    # ---- B1. 补缩进 ----
    set_first_line_indent(paras[i_kt])
    set_first_line_indent(paras[i_wt])

    # ---- B2/B3. Redis 重排 + 补图 ----
    img_win = new_img_para(d, IMG_WIN_REDIS)
    img_kylin = new_img_para(d, IMG_KYLIN_REDIS)
    # 依次插到「判定标准」之后：Windows 组在前、麒麟组在后（与前面各小节一致）
    chain = [paras[i_wt]._p, paras[i_w1]._p, img_win, paras[i_kt]._p,
             paras[i_a]._p, paras[i_b]._p, paras[i_c]._p]
    # K1~K4 依次取出（用文本重定位，因搬移后索引失效）
    k_elems = []
    for p in d.paragraphs:
        t = p.text.strip()
        if t.startswith(("核查日志保留策略", "核查持久化与日志文件位置", "核查慢查询日志", "核查日志开启情况")):
            k_elems.append(p._p)
    chain += k_elems + [img_kylin]

    cur = paras[i_stat]._p          # 判定标准段作锚
    for el in chain:
        if el is cur:
            continue
        cur.addnext(el)
        cur = el

    d.save(PATH)
    print("已保存", PATH)
    return 0


if __name__ == "__main__":
    sys.exit(main())
