# -*- coding: utf-8 -*-
"""《配置核查作业指导书》格式审计（只读，不修改文件）。

用法：python _audit_guide.py
输出：页面与页码、字体与字号、标题层级、段落格式、图片、判定标准覆盖、历史硬错误。
"""
import collections
import re
import sys
import zipfile

import docx
from lxml import etree
from docx.oxml.ns import qn

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
PATH = "配置核查作业指导书_v2.2.docx"
OK = "  [OK] "
WARN = "  [!!] "


def main():
    d = docx.Document(PATH)
    issues = []

    # ---------- 1. 页面与页码 ----------
    print("=== 1. 页面设置与页码 ===")
    print("  节数:", len(d.sections))
    s = d.sections[0]
    print("  纸张 %.1fx%.1fcm 边距 上%.2f/下%.2f/左%.2f/右%.2f" % (
        s.page_width.cm, s.page_height.cm, s.top_margin.cm,
        s.bottom_margin.cm, s.left_margin.cm, s.right_margin.cm))
    sx = re.sub(r"\s+", " ", s._sectPr.xml)
    pg = re.findall(r"<w:pgNumType[^>]*/?>", sx)
    print("  页码类型:", pg if pg else "无")
    print("  titlePg(首页不同):", "titlePg" in sx)
    fx = s.footer._element.xml
    instr = re.findall(r"<w:instrText[^>]*>([^<]*)</w:instrText>", fx)
    print("  页脚域:", [i.strip() for i in instr])
    if len(d.sections) < 3:
        issues.append("只有 %d 个节：标准的『封面无页码→目录罗马数字→正文阿拉伯从 1 起』三段式页码需要分节，当前未实现（现为全文连续阿拉伯数字，封面同样有页码）。" % len(d.sections))
    if "titlePg" not in sx:
        issues.append("未启用『首页不同』，封面上会显示页码。")

    # ---------- 2. 目录 ----------
    print("\n=== 2. 目录 ===")
    z = zipfile.ZipFile(PATH)
    root = etree.fromstring(z.read("word/document.xml"))
    body = root.find(W + "body")
    sdt = [el for el in body if etree.QName(el).localname == "sdt"]
    if sdt:
        txt = "".join(sdt[0].itertext())
        entries = re.findall(r"HYPERLINK \\l _Toc\d+ ([^H]*?) PAGEREF", txt)
        print("  目录为内容控件(sdt)，条目数约:", len(entries))
        print("  前 5 条:", [e.strip() for e in entries[:5]])
        if any(e.strip() in ("配置核查", "作业指导书", "目录") for e in entries):
            issues.append("目录缓存里含『配置核查 / 作业指导书 / 目录』等封面条目，与当前标题结构不一致（封面标题现为正文样式）——交付前需右键目录『更新域』重新生成。")
    else:
        issues.append("未找到目录内容控件。")

    # ---------- 3. 字体与字号 ----------
    print("\n=== 3. 字体与字号 ===")
    fonts = collections.Counter()
    sizes = collections.Counter()
    nofont_head = []
    nofont_body = []
    for p in d.paragraphs:
        for r in p.runs:
            if not r.text.strip():
                continue
            rPr = r._element.rPr
            rf = rPr.rFonts if rPr is not None else None
            a = rf.get(qn("w:ascii")) if rf is not None else None
            e = rf.get(qn("w:eastAsia")) if rf is not None else None
            sz = r.font.size.pt if r.font.size else None
            fonts[(a, e)] += 1
            sizes[sz] += 1
            if e is None:
                (nofont_head if p.style.name.startswith("Heading") else nofont_body).append(
                    (p.style.name, p.text.strip()[:44]))
    for k, v in fonts.most_common(8):
        print("   %s -> %d" % (k, v))
    print("   字号分布:", dict(sizes.most_common(10)))
    if nofont_head:
        issues.append("%d 个标题未设置中文字体（eastAsia），会回退为默认字体，与其他标题不一致，例：%s" % (
            len(nofont_head), nofont_head[0][1]))
    if nofont_body:
        issues.append("%d 处正文未设置中文字体，例：%s" % (len(nofont_body), nofont_body[0][1]))

    # ---------- 4. 标题层级 ----------
    print("\n=== 4. 标题层级 ===")
    heads = [(p.style.name, p.text.strip()) for p in d.paragraphs
             if p.style.name.startswith("Heading") and p.text.strip()]
    print("   合计 %d：%s" % (len(heads), dict(collections.Counter(h[0] for h in heads))))
    empty = sum(1 for p in d.paragraphs if p.style.name.startswith("Heading") and not p.text.strip())
    print("   空标题:", empty)
    if empty:
        issues.append("存在 %d 个空标题段落。" % empty)
    nonum = [(l, t) for l, t in heads if not re.match(r"^\d", t)]
    print("   非编号标题:", len(nonum), [t[:30] for _, t in nonum[:4]])
    if nonum:
        issues.append("有 %d 个标题不是编号形式，疑似正文被误设为标题样式，例：%s" % (len(nonum), nonum[0][1]))
    nums = [re.match(r"^(\d+(?:\.\d+)*)", t).group(1) for _, t in heads if re.match(r"^\d", t)]
    dup = [k for k, v in collections.Counter(nums).items() if v > 1]
    print("   重复编号:", len(dup))
    if dup:
        issues.append("重复编号：%s" % dup[:8])
    stack = {}
    mism = []
    for lvl, t in heads:
        m = re.match(r"^(\d+(?:\.\d+)*)", t)
        if not m:
            continue
        num = m.group(1)
        n = lvl.split()[-1]
        if n == "2":
            stack[2] = num
        elif n == "3":
            if stack.get(2) and not num.startswith(stack[2] + "."):
                mism.append((num, stack[2]))
            stack[3] = num
        elif n == "4":
            if stack.get(3) and not num.startswith(stack[3] + "."):
                mism.append((num, stack[3]))
    print("   层级前缀不匹配:", len(mism), mism[:5])
    if mism:
        issues.append("有 %d 处标题层级与父级编号不匹配，例：%s 挂在 %s 下" % (len(mism), mism[0][0], mism[0][1]))
    # 编号后空格
    sp = nosp = 0
    nosp_s = []
    for _, t in heads:
        m = re.match(r"^(\d+(?:\.\d+)+|\d+)(.*)$", t)
        if not m:
            continue
        if m.group(2)[:1] in (" ", "\u3000"):
            sp += 1
        else:
            nosp += 1
            if len(nosp_s) < 3:
                nosp_s.append(t[:34])
    print("   编号后带空格 %d / 不带空格 %d" % (sp, nosp))
    if nosp:
        issues.append("标题编号与标题文字之间的空格不统一：%d 个标题编号后无空格（如「%s」），其余 %d 个有空格。" % (nosp, nosp_s[0], sp))

    # ---------- 5. 段落与图片 ----------
    print("\n=== 5. 段落与图片 ===")
    agg = collections.defaultdict(lambda: collections.Counter())
    for p in d.paragraphs:
        pf = p.paragraph_format
        agg[p.style.name][(str(pf.alignment),
                           round(pf.first_line_indent.pt, 1) if pf.first_line_indent else None)] += 1
    for name in ("Normal", "HTML Preformatted", "Normal (Web)"):
        if name in agg:
            print("   %s: %s" % (name, dict(agg[name].most_common(3))))
    print("   内联图:", len(d.inline_shapes))
    al = collections.Counter()
    for p in d.paragraphs:
        if p._p.findall(".//" + qn("w:drawing")):
            al[str(p.paragraph_format.alignment)] += 1
    print("   图段对齐:", dict(al))
    if any(k != "CENTER (1)" for k in al):
        issues.append("存在非居中的图片段落。")
    pre = [p for p in d.paragraphs if p.style.name == "HTML Preformatted"]
    if pre:
        szs = sorted({r.font.size.pt if r.font.size else None for p in pre for r in p.runs})
        print("   代码块段落 %d，字号 %s" % (len(pre), szs))
        if szs and szs[0] and szs[0] < 9:
            issues.append("代码块字号仅 %spt（正文 14pt），打印与阅读都很吃力。" % szs[0])
    web = sum(1 for p in d.paragraphs if p.style.name == "Normal (Web)")
    if web:
        issues.append("有 %d 段使用网页样式『Normal (Web)』（粘贴带入），样式名与正文不统一。" % web)
    caps = [p.text for p in d.paragraphs if re.match(r"^图\s*\d+[-–]\d+", p.text.strip())]
    print("   题注残留:", len(caps))

    # ---------- 6. 判定标准与硬错误 ----------
    print("\n=== 6. 内容基线 ===")
    ch = h2 = None
    have = collections.defaultdict(lambda: [False])
    for p in d.paragraphs:
        t = p.text.strip()
        if p.style.name == "Heading 1":
            ch = t
        elif p.style.name == "Heading 2":
            h2 = t
            if ch and ch[:1] in ("1", "2"):
                have[h2] = [False]
        elif h2 and t.startswith("判定标准") and have.get(h2):
            have[h2][-1] = True
    c1 = [k for k in have if k.startswith("1.")]
    c2 = [k for k in have if k.startswith("2.")]
    print("   第1章判定标准 %d/%d，第2章 %d/%d" % (
        sum(1 for k in c1 if all(have[k])), len(c1), sum(1 for k in c2 if all(have[k])), len(c2)))
    txt = "\n".join(p.text for p in d.paragraphs)
    pats = ["SERVERPROPERTYC", r"ROUTINE_ TYPE", r"ROUTINE TYPE/NAME", "validate_password.dll",
            "employeesFOR", "触发器DELIMITER", "END IF;END"]
    hits = {p: len(re.findall(p, txt)) for p in pats}
    print("   历史硬错误:", hits)
    if any(hits.values()):
        issues.append("仍有历史硬错误未清零：%s" % {k: v for k, v in hits.items() if v})

    # ---------- 结论 ----------
    print("\n=== 审计结论：%d 项待处理 ===" % len(issues))
    for i, x in enumerate(issues, 1):
        print("  %d. %s" % (i, x))
    return 0


if __name__ == "__main__":
    sys.exit(main())
