# -*- coding: utf-8 -*-
r"""指导书封面加校徽（与操作说明书/交叉验证报告封面版式一致）。

封面原状：空行、空格段、空行、标题两行、8 空行、说明行。
改后：空行 → 校徽（居中 3.5cm）→ 空行 → 标题 → 8 空行 → 说明行。

用法：python _guide_badge.py
"""
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

PATH = "配置核查作业指导书_v2.0.0.docx"
BADGE = "测评报告/说明书截图/npu_badge.png"

d = docx.Document(PATH)
ps = d.paragraphs
assert ps[3].text.strip() == "配置核查", "封面结构异常：" + repr(ps[3].text[:20])

pic = d.add_paragraph()
pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
pic.paragraph_format.first_line_indent = None
pic.paragraph_format.space_after = Pt(18)
pic.add_run().add_picture(BADGE, width=Cm(3.5))
ps[1]._p.addprevious(pic._p)      # 插到第一个空行之后
ps[2]._p.getparent().remove(ps[2]._p)   # 去掉多余空段，保持标题位置
d.save(PATH)

chk = docx.Document(PATH)
from docx.oxml.ns import qn
for i, p in enumerate(chk.paragraphs[:6]):
    n = len(p._p.findall('.//' + qn('w:drawing')))
    print(i, repr(p.text[:16]), "图=%d" % n)
print("已保存")
